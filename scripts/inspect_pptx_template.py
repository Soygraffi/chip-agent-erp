# -*- coding: utf-8 -*-
"""Dump PPTX template structure for replicating in a Word quote."""
import json
import sys
from pptx import Presentation
from pptx.util import Emu

SRC = r"C:\Users\pengh\Desktop\半导体元器件代理报价单模板A4横版3.pptx"

def emu_cm(v):
    try:
        return round(Emu(v).cm, 2)
    except Exception:
        return None

def dump_shape(sh, indent=0):
    out = {
        "type": str(sh.shape_type),
        "name": sh.name,
        "left_cm": emu_cm(sh.left),
        "top_cm": emu_cm(sh.top),
        "width_cm": emu_cm(sh.width),
        "height_cm": emu_cm(sh.height),
    }
    if sh.has_text_frame:
        out["text"] = sh.text_frame.text
        out["paras"] = []
        for p in sh.text_frame.paragraphs:
            runs = []
            for r in p.runs:
                runs.append({
                    "text": r.text,
                    "size_pt": r.font.size.pt if r.font.size else None,
                    "bold": r.font.bold,
                    "color": str(r.font.color.rgb) if r.font.color and r.font.color.type is not None else None,
                    "font": r.font.name,
                })
            out["paras"].append({
                "align": str(p.alignment),
                "text": p.text,
                "runs": runs,
            })
    if getattr(sh, "has_table", False) and sh.has_table:
        tbl = sh.table
        out["table"] = {
            "rows": len(tbl.rows),
            "cols": len(tbl.columns),
            "col_widths_cm": [emu_cm(c.width) for c in tbl.columns],
            "row_heights_cm": [emu_cm(r.height) for r in tbl.rows],
            "cells": [
                [cell.text for cell in row.cells]
                for row in tbl.rows
            ],
        }
    return out

prs = Presentation(SRC)
result = {
    "slide_width_cm": emu_cm(prs.slide_width),
    "slide_height_cm": emu_cm(prs.slide_height),
    "slides": [],
}
for idx, slide in enumerate(prs.slides):
    s = {"index": idx + 1, "shapes": []}
    for sh in slide.shapes:
        try:
            s["shapes"].append(dump_shape(sh))
        except Exception as e:
            s["shapes"].append({"name": sh.name, "error": str(e)})
    result["slides"].append(s)

print(json.dumps(result, ensure_ascii=False, indent=1))
